"""Statements and spreadsheets as files: Excel and PDF.

What matters is that the files say the same as the screen and the books: a
statement's closing balance is what the borrower owes, the member register has
each member once with the balances the loan and savings books hold, and every
report that offered CSV offers Excel with the same rows.
"""
import io
import os
import tempfile
from datetime import date
from decimal import Decimal

from django.db import connection
from django.test import override_settings
from django.test.utils import CaptureQueriesContext

from core.models import BorrowerGroup, GroupMember, Loan, SavingsProduct, User
from core.services import savings as savings_svc
from core.services import spreadsheets as sheets

from .test_arrears import ArrearsParityBase
from .test_components import LedgerBase, dec


def workbook(response):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(response.content))


def table(ws):
    """A sheet's table as dicts, found by its header row (the one with a fill)."""
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, row in enumerate(rows) if row and ws.cell(row=i + 1, column=1).fill.fgColor
                 and ws.cell(row=i + 1, column=1).fill.fill_type == "solid")
    header = rows[start]
    out = []
    for row in rows[start + 1:]:
        if row[0] == "Total" or all(v is None for v in row):
            continue
        out.append(dict(zip(header, row)))
    return out


class DownloadBase(LedgerBase):
    def setUp(self):
        super().setUp()
        self.product = self.make_product()
        self.borrower = self.make_borrower()
        self.loan = self.disbursed_loan(self.product, self.borrower)  # 1000 at 5% over 6

    def pay(self, amount, on="2026-03-25", reference=None):
        response = self.teller.post(f"/api/loans/{self.loan['id']}/repayments", {
            "amount": amount, "method": "bank_transfer", "txn_date": on,
            "reference": reference}, format="json")
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()


class LoanStatementTests(DownloadBase):
    def test_the_closing_balance_is_what_the_borrower_owes_that_has_been_charged(self):
        self.pay("197.02", reference="PAY-1")
        self.officer.post(f"/api/loans/{self.loan['id']}/accrue-penalties?as_of=2026-05-10")
        self.officer.post(f"/api/loans/{self.loan['id']}/charges", {
            "name": "Statement reprint", "amount": "3.00", "collection": "balance"},
            format="json")
        wrong = self.pay("50.00", on="2026-05-11", reference="PAY-X")
        self.officer.post(f"/api/loans/{self.loan['id']}/transactions/{wrong['id']}/reverse",
                          {"narration": "Wrong loan"}, format="json")

        body = self.officer.get(f"/api/loans/{self.loan['id']}/statement").json()
        loan = Loan.objects.get(pk=self.loan["id"])
        owed = loan.principal_outstanding + loan.penalties_outstanding + loan.charges_outstanding
        self.assertEqual(dec(body["closing_balance"]), owed)
        self.assertEqual(dec(body["lines"][-1]["balance"]), owed)
        # The reversed receipt and its reversal are both on it, so the borrower's
        # own receipt still appears.
        kinds = [line["type"] for line in body["lines"]]
        self.assertIn("reversal", kinds)
        self.assertTrue(any(l["reversed"] for l in body["lines"]))

    def test_interest_paid_has_its_own_column_and_does_not_touch_the_balance(self):
        self.pay("197.02")
        body = self.officer.get(f"/api/loans/{self.loan['id']}/statement").json()
        repayment = next(l for l in body["lines"] if l["type"] == "repayment")
        self.assertEqual(dec(repayment["interest"]), dec("50.00"))
        self.assertEqual(dec(repayment["balance"]), dec("1000.00") - dec(repayment["principal"]))

    def test_a_period_carries_the_balance_forward(self):
        self.pay("197.02", on="2026-03-25")
        self.pay("197.02", on="2026-04-25")
        body = self.officer.get(
            f"/api/loans/{self.loan['id']}/statement?start=2026-04-01").json()
        self.assertEqual(len(body["lines"]), 1)
        self.assertGreater(dec(body["opening_balance"]), dec("800"))
        self.assertEqual(dec(body["closing_balance"]),
                         Loan.objects.get(pk=self.loan["id"]).principal_outstanding)

    def test_the_excel_statement_has_the_lines_and_the_schedule(self):
        self.pay("197.02", reference="PAY-1")
        response = self.officer.get(f"/api/loans/{self.loan['id']}/statement?fmt=xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(f"statement_{self.loan['loan_no']}.xlsx", response["Content-Disposition"])
        book = workbook(response)
        self.assertEqual(book.sheetnames, ["Statement", "Schedule"])
        lines = table(book["Statement"])
        self.assertEqual(lines[-1]["Reference"], "PAY-1")
        self.assertEqual(Decimal(str(lines[-1]["Paid"])), dec("197.02"))
        self.assertEqual(len(table(book["Schedule"])), 6)

    def test_the_pdf_statement_is_a_pdf(self):
        self.pay("197.02")
        response = self.officer.get(f"/api/loans/{self.loan['id']}/statement?fmt=pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn(f"statement_{self.loan['loan_no']}.pdf", response["Content-Disposition"])

    def test_the_pdf_has_no_logo_unless_one_is_configured(self):
        path = f"/api/loans/{self.loan['id']}/statement?fmt=pdf"
        # Pinned blank: what this machine's .env names must not decide the result.
        with override_settings(STATEMENT_LOGO=""):
            self.assertNotIn(b"/Subtype /Image", self.officer.get(path).content)
        with override_settings(STATEMENT_LOGO="no/such/file.png"):
            self.assertNotIn(b"/Subtype /Image", self.officer.get(path).content)

        from PIL import Image

        with tempfile.TemporaryDirectory() as folder:
            logo = os.path.join(folder, "logo.png")
            Image.new("RGB", (40, 20), "navy").save(logo)
            with override_settings(STATEMENT_LOGO=logo):
                self.assertIn(b"/Subtype /Image", self.officer.get(path).content)

    def test_the_logo_that_ships_with_the_code_is_found_from_the_backend_directory(self):
        from core import documents

        with override_settings(STATEMENT_LOGO="branding/zinmad-mark.png"):
            self.assertIsNotNone(documents.logo_path(), "backend/branding/zinmad-mark.png is missing")
            pdf = self.officer.get(f"/api/loans/{self.loan['id']}/statement?fmt=pdf").content
            self.assertIn(b"/Subtype /Image", pdf)
            agreement = self.officer.get(f"/api/loans/{self.loan['id']}/agreement").content.decode()
            self.assertIn('src="data:image/png;base64,', agreement)
        with override_settings(STATEMENT_LOGO=""):
            agreement = self.officer.get(f"/api/loans/{self.loan['id']}/agreement").content.decode()
            self.assertNotIn("<img", agreement)


class SavingsStatementTests(DownloadBase):
    def setUp(self):
        super().setUp()
        teller = User.objects.get(username="teller")
        product = SavingsProduct.objects.create(code="SAV", name="Ordinary savings")
        from core.models import Borrower

        self.account = savings_svc.open_account(Borrower.objects.get(pk=self.borrower["id"]),
                                                product, teller)
        savings_svc.deposit(self.account, teller, dec("300"), date(2026, 3, 1), "cash", "D1")
        out = savings_svc.withdraw(self.account, teller, dec("120"), date(2026, 3, 5), "cash")
        savings_svc.reverse(self.account, out, teller, "Paid to the wrong member")
        savings_svc.deposit(self.account, teller, dec("50"), date(2026, 3, 9), "cash", "D2")

    def test_the_closing_balance_is_the_account_balance(self):
        body = self.officer.get(f"/api/savings/accounts/{self.account.id}/statement").json()
        self.account.refresh_from_db()
        self.assertEqual(dec(body["closing_balance"]), self.account.balance)
        self.assertEqual(dec(body["total_in"]), dec("470.00"))  # 300 + reversed 120 + 50
        self.assertEqual(dec(body["total_out"]), dec("120.00"))

    def test_excel_and_pdf(self):
        xlsx = self.officer.get(f"/api/savings/accounts/{self.account.id}/statement?fmt=xlsx")
        lines = table(workbook(xlsx)["Statement"])
        self.assertEqual([l["Reference"] for l in lines if l["Reference"]], ["D1", "D2"])
        self.account.refresh_from_db()
        self.assertEqual(Decimal(str(lines[-1]["Balance"])), self.account.balance)
        pdf = self.officer.get(f"/api/savings/accounts/{self.account.id}/statement?fmt=pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))


class ExcelEverywhereTests(DownloadBase):
    """Every listing that offered CSV offers Excel, with the same rows."""

    ENDPOINTS = ["/api/reports/loan-book", "/api/reports/par?as_of=2026-09-01",
                 "/api/reports/transactions?start=2026-01-01&end=2026-12-31",
                 "/api/ledger/trial-balance", "/api/ledger/journal", "/api/savings/accounts",
                 "/api/reports/audit", "/api/journals", "/api/tills",
                 "/api/funding/capital", "/api/groups/performance"]

    def test_each_listing_downloads_as_excel(self):
        for path in self.ENDPOINTS:
            joiner = "&" if "?" in path else "?"
            response = self.admin.get(f"{path}{joiner}fmt=xlsx")
            self.assertEqual(response.status_code, 200, f"{path}: {response.content[:200]}")
            self.assertIn("spreadsheetml", response["Content-Type"], path)
            workbook(response)  # opens

    def test_the_excel_rows_match_the_csv_rows(self):
        csv_text = self.admin.get("/api/reports/loan-book?fmt=csv").content.decode()
        csv_rows = csv_text.strip().splitlines()[1:]
        sheet = workbook(self.admin.get("/api/reports/loan-book?fmt=xlsx")).active
        rows = table(sheet)
        self.assertEqual(len(rows), len(csv_rows))
        self.assertEqual(rows[0]["Loan no"], self.loan["loan_no"])
        # A number Excel can sum, not text (a whole amount reads back as an int).
        self.assertIsInstance(rows[0]["Principal outstanding"], (int, float))


class MemberRegisterTests(DownloadBase):
    def setUp(self):
        super().setUp()
        self.other = self.make_borrower(national_id="63-222222B63", first_name="Second")
        teller = User.objects.get(username="teller")
        from core.models import Borrower

        member = Borrower.objects.get(pk=self.borrower["id"])
        # One account per product per member, so two products for two accounts.
        for code, amount in (("SAV", "100"), ("FIX", "250")):
            product = SavingsProduct.objects.create(code=code, name=f"Savings {code}")
            account = savings_svc.open_account(member, product, teller)
            savings_svc.deposit(account, teller, dec(amount), date(2026, 3, 1), "cash")
        group = BorrowerGroup.objects.create(group_no="GRP-1", name="Tashinga")
        GroupMember.objects.create(group=group, borrower=member, role="treasurer")

    def test_each_member_appears_once_with_their_balances(self):
        rows = {r["member_no"]: r for r in sheets.member_register()}
        self.assertEqual(len(rows), 2)
        mine = rows[self.borrower["borrower_no"]]
        loan = Loan.objects.get(pk=self.loan["id"])
        self.assertEqual(mine["savings_accounts"], 2)
        self.assertEqual(mine["savings_balance"], dec("350.00"))
        self.assertEqual(mine["active_loans"], 1)
        self.assertEqual(mine["total_outstanding"], loan.total_outstanding)
        self.assertEqual(mine["group"], "Tashinga")
        self.assertEqual(mine["group_role"], "Treasurer")
        # Disbursed in March and never paid: overdue, counted from the first due date.
        self.assertGreater(mine["overdue_amount"], 0)
        self.assertEqual(mine["days_overdue"], (date.today() - date(2026, 3, 25)).days)
        nobody = rows[self.other["borrower_no"]]
        self.assertEqual(nobody["total_outstanding"], dec("0"))
        self.assertEqual(nobody["savings_balance"], dec("0.00"))

    def test_the_register_downloads_with_a_total_under_the_money(self):
        response = self.admin.get("/api/reports/spreadsheets/members?fmt=xlsx")
        self.assertEqual(response.status_code, 200)
        ws = workbook(response).active
        self.assertEqual(len(table(ws)), 2)
        formulas = [c.value for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=SUBTOTAL")]
        self.assertTrue(formulas, "no subtotal row under the money columns")

    def test_every_spreadsheet_answers_json_csv_and_excel(self):
        for kind in sheets.SPREADSHEETS:
            base = f"/api/reports/spreadsheets/{kind}"
            self.assertEqual(self.viewer().get(base).status_code, 200, kind)
            self.assertTrue(self.admin.get(f"{base}?fmt=csv")["Content-Type"]
                            .startswith("text/csv"), kind)
            workbook(self.admin.get(f"{base}?fmt=xlsx"))

    def test_the_portfolio_workbook_has_a_sheet_for_each(self):
        response = self.admin.get("/api/reports/workbook")
        self.assertEqual(response.status_code, 200)
        book = workbook(response)
        self.assertEqual(book.sheetnames, ["Summary", "Members", "Loans outstanding", "Overdue",
                                           "Savings balances", "Group membership"])
        self.assertEqual(len(table(book["Group membership"])), 1)

    def test_the_json_names_its_money_columns(self):
        # JSON renders a Decimal as a float, so the screen is told which are money.
        data = self.admin.get("/api/reports/spreadsheets/members").json()
        self.assertEqual(data["count"], 2)
        self.assertIn("total_outstanding", data["money_columns"])
        self.assertIn("savings_balance", data["money_columns"])
        self.assertNotIn("days_overdue", data["money_columns"])
        self.assertNotIn("member_no", data["money_columns"])

    def test_the_screen_gets_the_ids_to_open_each_row_and_the_files_do_not(self):
        rows = self.admin.get("/api/reports/spreadsheets/members").json()["results"]
        self.assertTrue(all(row["borrower_id"] for row in rows))
        loans = self.admin.get("/api/reports/spreadsheets/loans-outstanding").json()["results"]
        self.assertTrue(all(row["loan_id"] and row["borrower_id"] for row in loans))
        csv = self.admin.get("/api/reports/spreadsheets/members?fmt=csv").content.decode()
        self.assertNotIn("borrower_id", csv.splitlines()[0])

    def test_a_search_narrows_the_rows_across_every_page(self):
        rows = self.admin.get("/api/reports/spreadsheets/members").json()["results"]
        wanted = rows[0]["member_no"]
        found = self.admin.get(f"/api/reports/spreadsheets/members?search={wanted.lower()}").json()
        self.assertEqual([r["member_no"] for r in found["results"]], [wanted])
        self.assertEqual(found["total_rows"], 2)

    def test_an_unknown_spreadsheet_is_refused(self):
        self.assertEqual(self.admin.get("/api/reports/spreadsheets/nonsense").status_code, 400)

    def viewer(self):
        from core.models import Role

        if not User.objects.filter(username="viewer").exists():
            User.objects.create_user("viewer", "viewer123", full_name="Viewer", role=Role.USER)
        return self.client_for("viewer", "viewer123")


class SpreadsheetQueryCountTests(ArrearsParityBase):
    """The register reads the book in a fixed number of queries, however many members."""

    def counts_for(self, call):
        with CaptureQueriesContext(connection) as first:
            call()
        before = len(first)
        for index in range(7, 14):
            borrower = self.make_borrower(national_id=f"63-6000{index:02d}A63")
            self.disbursed_loan(self.product, borrower, principal=1500,
                                disbursement_date="2026-02-01")
        with CaptureQueriesContext(connection) as second:
            call()
        return before, len(second)

    def test_the_member_register_does_not_query_per_member(self):
        self.busy_book()
        before, after = self.counts_for(lambda: sheets.member_register())
        self.assertEqual(before, after, f"member register queries grew {before} -> {after}")

    def test_the_workbook_does_not_query_per_member(self):
        self.busy_book()
        before, after = self.counts_for(lambda: sheets.portfolio_workbook())
        self.assertEqual(before, after, f"workbook queries grew {before} -> {after}")
